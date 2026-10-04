from __future__ import annotations

import json
from pathlib import Path
import zipfile

from docx import Document
from openpyxl import load_workbook

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_PREVIEW_CHARS = 24000


class OfficeFileError(ValueError):
    pass


def inspect_office(path: Path) -> str:
    if path.suffix.lower() not in {".xlsx", ".docx"}:
        raise OfficeFileError("Hozircha .xlsx va .docx fayllarini yuboring.")
    if path.stat().st_size > MAX_UPLOAD_BYTES:
        raise OfficeFileError("Fayl hajmi 10 MB dan oshmasin.")
    try:
        with zipfile.ZipFile(path) as archive:
            items = archive.infolist()
            if len(items) > 3000 or sum(item.file_size for item in items) > 40 * 1024 * 1024:
                raise OfficeFileError("Fayl ochilganda juda katta bo‘ladi. Kichikroq nusxa yuboring.")
            if any("vbaProject" in item.filename for item in items):
                raise OfficeFileError("Makrosli hujjatlarni qabul qilmayman.")
        lines = []
        used = 0
        truncated = False
        def add(value):
            nonlocal used, truncated
            line = str(value)
            remaining = MAX_PREVIEW_CHARS - used
            if remaining <= 0:
                truncated = True
                return False
            lines.append(line[:remaining])
            used += len(line[:remaining]) + 1
            if len(line) > remaining:
                truncated = True
                return False
            return True
        if path.suffix.lower() == ".xlsx":
            book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                if len(book.worksheets) > 5:
                    truncated = True
                for ws in book.worksheets[:5]:
                    add(f"Sheet: {ws.title}; dimension hint: {ws.max_row} rows, {ws.max_column} columns")
                    # Do not trust worksheet dimensions; stream with explicit bounds.
                    ws.reset_dimensions()
                    for index, row in enumerate(ws.iter_rows(max_col=31, values_only=True)):
                        if index >= 200:
                            truncated = True
                            break
                        if row[30] is not None:
                            truncated = True
                        if not add(json.dumps({"row": index + 1, "cells": list(row[:30])}, ensure_ascii=False, default=str)):
                            break
                    if used >= MAX_PREVIEW_CHARS:
                        break
            finally:
                book.close()
        else:
            doc = Document(path)
            for paragraph in doc.paragraphs:
                if not add(paragraph.text):
                    break
            if used < MAX_PREVIEW_CHARS:
                for table in doc.tables:
                    if not add("Table:"):
                        break
                    for row in table.rows:
                        if not add(json.dumps([cell.text for cell in row.cells], ensure_ascii=False)):
                            break
                    if used >= MAX_PREVIEW_CHARS:
                        break
        note = (
            "PREVIEW IS TRUNCATED: analyse only the visible data; request a smaller file for full processing."
            if truncated else
            "This preview contains extracted text/cells only, not formatting, embedded images or charts."
        )
        return note + "\n" + "\n".join(lines)
    except OfficeFileError:
        raise
    except Exception as exc:
        raise OfficeFileError("Faylni o‘qib bo‘lmadi. To‘g‘ri .xlsx yoki .docx nusxa yuboring.") from exc
