from __future__ import annotations

import ast
import math
import operator
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.properties import CalcProperties

class WorkToolError(ValueError):
    pass


def calculate(expression: str) -> str:
    """Evaluate a bounded arithmetic AST, never Python code."""
    if not isinstance(expression, str) or len(expression) > 500:
        raise WorkToolError("Expression too long")
    tree = ast.parse(expression, mode="eval")
    if sum(1 for _ in ast.walk(tree)) > 100:
        raise WorkToolError("Expression too complex")
    operations = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul,
                  ast.Div: operator.truediv, ast.Mod: operator.mod, ast.Pow: operator.pow}
    functions = {"sqrt": math.sqrt, "sin": math.sin, "cos": math.cos, "tan": math.tan,
                 "log": math.log, "log10": math.log10, "exp": math.exp, "abs": abs,
                 "round": round}
    def checked(value):
        if type(value) not in {int, float} or not math.isfinite(value) or abs(value) > 1e100:
            raise WorkToolError("Result outside calculator range")
        return value
    def visit(node):
        if isinstance(node, ast.Constant):
            return checked(node.value)
        if isinstance(node, ast.Name) and node.id in {"pi", "e"}:
            return getattr(math, node.id)
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            value = visit(node.operand)
            return checked(value if isinstance(node.op, ast.UAdd) else -value)
        if isinstance(node, ast.BinOp) and type(node.op) in operations:
            left, right = visit(node.left), visit(node.right)
            if isinstance(node.op, ast.Pow) and abs(right) > 100:
                raise WorkToolError("Exponent too large")
            return checked(operations[type(node.op)](left, right))
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id in functions:
            if node.keywords or not 1 <= len(node.args) <= 2:
                raise WorkToolError("Invalid function arguments")
            return checked(functions[node.func.id](*(visit(arg) for arg in node.args)))
        raise WorkToolError("Only arithmetic and listed math functions are allowed")
    return str(visit(tree.body))


def _text(value, limit=8000):
    if not isinstance(value, str) or len(value) > limit:
        raise WorkToolError("Text field is invalid or too long")
    return value


def _rows(rows, max_rows=500):
    if not isinstance(rows, list) or len(rows) > max_rows:
        raise WorkToolError("Too many rows")
    width = max((len(row) for row in rows if isinstance(row, list)), default=0)
    if width > 30 or any(not isinstance(row, list) for row in rows):
        raise WorkToolError("Invalid table")
    for row in rows:
        for cell in row:
            if cell is not None and type(cell) not in {str, int, float, bool}:
                raise WorkToolError("Invalid cell value")
            if isinstance(cell, str):
                _text(cell, 2000)
            elif isinstance(cell, float) and not math.isfinite(cell):
                raise WorkToolError("Invalid number")
    return width


def _formula(value: str) -> None:
    # Restricted to local A1 references, numeric arithmetic and a small function set.
    import re
    if len(value) > 500 or re.search(r'[^A-Za-z0-9_=+*/^().,: $<>-]', value):
        raise WorkToolError("Unsupported formula")
    names = re.findall(r'([A-Za-z_][A-Za-z_0-9]*)\s*\(', value)
    if any(name.upper() not in {"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "ROUND", "ABS", "IF"} for name in names):
        raise WorkToolError("Unsupported Excel function")
    tokens = re.findall(r'[A-Za-z_][A-Za-z_0-9]*', value.replace('$', ''))
    if any(token.upper() not in {"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "ROUND", "ABS", "IF"}
           and not re.fullmatch(r'[A-Za-z]{1,3}[1-9][0-9]{0,6}', token) for token in tokens):
        raise WorkToolError("Only local cell references are supported")


def create_excel(data: dict, directory: Path, index: int) -> Path:
    sheets = data.get("sheets")
    if not isinstance(sheets, list) or not 1 <= len(sheets) <= 5:
        raise WorkToolError("Use 1 to 5 sheets")
    book = Workbook()
    book.remove(book.active)
    book.calculation = CalcProperties(fullCalcOnLoad=True)
    for sheet in sheets:
        name = _text(sheet["name"], 31)
        if not name or any(char in name for char in "[]:*?/\\") or name.casefold() in {n.casefold() for n in book.sheetnames}:
            raise WorkToolError("Invalid or duplicate sheet name")
        rows = sheet["rows"]
        width = _rows(rows)
        ws = book.create_sheet(name)
        for row in rows:
            ws.append(row)
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    if cell.value.startswith("="):
                        _formula(cell.value)
                    else:
                        cell.data_type = "s"
                if type(cell.value) in {int, float}:
                    cell.number_format = "#,##0.00"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if rows and width:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="18344A")
            from openpyxl.utils import get_column_letter
            for col in range(1, width + 1):
                ws.column_dimensions[get_column_letter(col)].width = 22
    path = directory / f"shadow-workbook-{index}.xlsx"
    book.save(path)
    book.close()
    return path


def create_word(data: dict, directory: Path, index: int) -> Path:
    title = _text(data["title"], 300)
    sections = data["sections"]
    if not isinstance(sections, list) or len(sections) > 40:
        raise WorkToolError("Too many sections")
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.8)
    doc.add_heading(title, 0)
    for section in sections:
        heading = _text(section["heading"], 300)
        if heading:
            doc.add_heading(heading, 1)
        paragraphs = section["paragraphs"]
        if not isinstance(paragraphs, list) or len(paragraphs) > 60:
            raise WorkToolError("Too many paragraphs")
        for paragraph in paragraphs:
            doc.add_paragraph(_text(paragraph))
        rows = section["table"]
        width = _rows(rows, 100)
        if rows and width:
            table = doc.add_table(rows=0, cols=width)
            table.style = "Light Shading Accent 1"
            for row in rows:
                cells = table.add_row().cells
                for index_cell, value in enumerate(row):
                    cells[index_cell].text = "" if value is None else str(value)
            for cell in table.rows[0].cells:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    path = directory / f"shadow-document-{index}.docx"
    doc.save(path)
    return path
