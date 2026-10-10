import asyncio
import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import httpx
from docx import Document
from openpyxl import load_workbook

from shadow.assistant import ShadowAssistant
from shadow.config import AIProvider
from shadow.development import DevelopmentStudio
from shadow.workspace import ChatStore, Workspace, WorkspaceError, workspace_reply
from shadow.workspace_agent import AgentRunner, compat_agent, MAX_CALLS
from tests.test_ai_fallback import make_settings, build, QuotaError, BadRequest
from tests.test_compat_tools import tool_call, completion


class WorkspaceAgentTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.worker = Workspace(ChatStore(Path(self.tmp.name)/'private.sqlite3'))
        self.assistant, self.primary, self.compat = build(make_settings(ai_work_mode='free', ai_base_url='https://openrouter.ai/api/v1', ai_model='one:free'))

    async def asyncTearDown(self):
        await self.worker.close()

    def runner(self, development=None):
        return AgentRunner(self.assistant, Path(self.tmp.name)/'artifacts', 11, {}, development)

    async def test_model_uses_tool_result_then_creates_a_downloadable_private_file(self):
        self.compat.chat.completions.create.side_effect = [
            completion(calls=[tool_call('a', 'calculate', {'expression': '125*8'})]),
            completion(calls=[tool_call('b', 'create_file', {'name': 'budget.csv', 'content': 'total\n1000\n'})]),
            completion('Budget tayyor, jami 1000.')]
        job = self.worker.start(11, self.assistant, 'Hisoblab CSV fayl qilib ber', 'agent')
        await self.worker.jobs[job['id']]['task']
        finished = self.worker.jobs[job['id']]
        self.assertEqual(finished['state'], 'done')
        thread = self.worker.store.get(11, job['thread_id'])
        answer = thread['messages'][-1]
        self.assertEqual([e['state'] for e in answer['events']], ['done', 'done'])
        path, item = self.worker.attachment(11, thread['id'], answer['attachments'][0]['id'])
        self.assertEqual(path.read_text(), 'total\n1000\n')
        self.assertEqual(item['name'], 'budget.csv')
        self.assertEqual(path.stat().st_mode & 0o777, 0o600)
        calls = self.compat.chat.completions.create.await_args_list
        self.assertEqual(json.loads(calls[1].kwargs['messages'][-1]['content'])['result'], '1000')
        self.assertTrue(any(t['function']['name']=='create_file' for t in calls[0].kwargs['tools']))
        with self.assertRaises(WorkspaceError):
            self.worker.attachment(22, thread['id'], item['id'])
        self.worker.delete(11, thread['id'])
        self.assertFalse(path.exists())

    async def test_openai_responses_uses_native_tool_calls_and_preserves_raw_code(self):
        assistant, primary, compat = build(make_settings())
        runner = AgentRunner(assistant, Path(self.tmp.name)/'outputs', 11, {})
        call = SimpleNamespace(type='function_call', call_id='call-1', name='create_file', arguments=json.dumps({'name': 'hello.py', 'content': "print('hello-world')\n"}))
        primary.responses.create.side_effect = [SimpleNamespace(output=[call], output_text=''), SimpleNamespace(output=[], output_text='Fayl tayyor')]
        answer, chosen = await workspace_reply(assistant, [{'role': 'user', 'content': 'Create hello.py'}], 'agent', [], runner=runner)
        self.assertEqual(answer, 'Fayl tayyor')
        self.assertEqual(chosen['provider'], 'OpenAI')
        self.assertFalse(primary.responses.create.await_args.kwargs['store'])
        self.assertEqual(primary.responses.create.await_args.kwargs['input'][-1]['type'], 'function_call_output')
        self.assertEqual((runner.directory/runner.files[0]['id']).read_text(), "print('hello-world')\n")
        compat.chat.completions.create.assert_not_awaited()

    async def test_word_and_excel_are_real_files(self):
        runner = self.runner()
        args = {'title': 'Report', 'sections': [{'heading': 'Numbers', 'paragraphs': ['Total: 1000'], 'table': []}]}
        result = json.loads(await runner.run('create_word', json.dumps(args)))
        self.assertTrue(result['ready_to_download'])
        doc = Document(runner.directory/runner.files[0]['id'])
        self.assertIn('Report', '\n'.join(p.text for p in doc.paragraphs))
        args = {'sheets': [{'name': 'Budget', 'rows': [['Item', 'Cost'], ['Hosting', 10]]}]}
        await runner.run('create_excel', json.dumps(args))
        with (runner.directory/runner.files[1]['id']).open('rb') as stream:
            wb = load_workbook(stream)
            self.assertEqual(wb['Budget']['B2'].value, 10)
            wb.close()

    async def test_syntax_check_does_not_execute_code_and_can_be_corrected(self):
        runner = self.runner()
        invalid = json.loads(await runner.run('validate_code', json.dumps({'language': 'python', 'content': 'def broken('})))
        self.assertFalse(invalid['valid'])
        self.assertEqual(invalid['line'], 1)
        valid = json.loads(await runner.run('validate_code', json.dumps({'language': 'python', 'content': "raise RuntimeError('do not execute')"})))
        self.assertTrue(valid['valid'])
        self.assertFalse(valid['execution'])

    async def test_search_uses_real_tool_sources_and_missing_connections_are_honest(self):
        runner = self.runner()
        results = {'results': [{'title': 'Source', 'url': 'https://example.com/', 'snippet': 'Evidence'}]}
        with patch('shadow.workspace_agent.search_plugin.run', AsyncMock(return_value=results)) as search:
            answer = json.loads(await runner.run('search_web', '{"query":"today","limit":3}'))
        self.assertEqual(answer, results)
        search.assert_awaited_once_with('today', 3)
        self.assertEqual(runner.sources, results['results'])
        with patch('shadow.workspace_agent.search_plugin.run', AsyncMock(side_effect=ValueError('not configured'))):
            failure = json.loads(await runner.run('search_web', '{"query":"other","limit":3}'))
        self.assertIn('error', failure)
        self.assertEqual(runner.events[-1]['state'], 'error')

    async def test_identical_actions_are_not_repeated_across_fallback(self):
        settings = make_settings(ai_work_mode='free', ai_base_url='https://openrouter.ai/api/v1', ai_model='one:free',
            ai_extra_providers=(AIProvider('second','https://openrouter.ai/api/v1','own-key','two:free',2),))
        args = {'name': 'result.txt', 'content': 'private output'}
        first = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=AsyncMock(side_effect=[completion(calls=[tool_call('a','create_file',args)]), QuotaError()]))))
        second = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=AsyncMock(side_effect=[completion(calls=[tool_call('b','create_file',args)]), completion('Tayyor')]))))
        with patch('shadow.assistant.AsyncOpenAI', side_effect=[first,second]):
            assistant = ShadowAssistant(settings)
        runner = AgentRunner(assistant,Path(self.tmp.name)/'outputs',11,{})
        answer, chosen = await workspace_reply(assistant,[{'role':'user','content':'create file'}],'agent',[],model='slot:1',runner=runner)
        self.assertEqual(answer,'Tayyor')
        self.assertTrue(chosen['fallback_used'])
        self.assertEqual(len(runner.files),1)
        self.assertEqual(len(runner.events),1)
        self.assertIn('Previously executed tool receipts',second.chat.completions.create.await_args_list[0].kwargs['messages'][-1]['content'])

    async def test_unknown_tools_paths_and_budgets_cannot_execute_commands(self):
        runner = self.runner()
        self.assertIn('error',json.loads(await runner.run('shell','{"command":"cat /etc/passwd"}')))
        self.assertIn('error',json.loads(await runner.run('create_file','{"name":"../private","content":"x"}')))
        self.assertEqual(runner.files,[])
        for i in range(MAX_CALLS):
            await runner.run('calculate',json.dumps({'expression':str(i)}))
        self.assertLessEqual(len(runner.events),MAX_CALLS)
        self.assertIn('error',json.loads(await runner.run('calculate','{"expression":"999"}')))
        self.assertIn('error',json.loads(await runner.run('calculate','[]')))

    async def test_completed_artifacts_remain_after_model_failure_or_cancel(self):
        for cancel in (False,True):
            async def failing(*args,runner,**kwargs):
                await runner.run('create_file','{"name":"saved.txt","content":"finished"}')
                if cancel:
                    await asyncio.Event().wait()
                raise RuntimeError('provider-private-secret')
            with patch('shadow.workspace.workspace_reply',side_effect=failing):
                job=self.worker.start(11,self.assistant,'Save output then stop','agent')
                running=self.worker.jobs[job['id']]['task']
                await asyncio.sleep(.03)
                if cancel:
                    await self.worker.cancel(11,job['id'])
                else:
                    await running
            thread=self.worker.store.get(11,job['thread_id'])
            self.assertEqual(len(thread['messages'][-1]['attachments']),1)
            item=thread['messages'][-1]['attachments'][0]
            path,_=self.worker.attachment(11,thread['id'],item['id'])
            self.assertEqual(path.read_text(),'finished')
            self.assertNotIn('provider-private-secret',json.dumps(self.worker.public(self.worker.jobs[job['id']])))

    async def test_project_worker_starts_and_is_partitioned_by_account(self):
        source=Path(self.tmp.name)/'shadow';source.mkdir();(source/'app.py').write_text("print('old')\n")
        proposal={'title':'Fix interface','summary':'Changed source','findings':[],'verification':['Check source'], 'changes':[{'path':'shadow/app.py','find':"print('old')",'replace':"print('new')"}]}
        generate=AsyncMock(side_effect=[{'files':['shadow/app.py']},proposal,proposal])
        dev=DevelopmentStudio(root=Path(self.tmp.name),state_dir=Path(self.tmp.name)/'dev',generate=generate)
        runner=self.runner(dev)
        result=json.loads(await runner.run('start_project_task','{"objective":"Fix the Shadow interface","kind":"build"}'))
        await dev.task
        self.assertEqual(result['state'],'queued')
        task=dev.get(result['id'],runner.scope)
        self.assertEqual(task['account_scope'],hashlib.sha256(b'11').hexdigest()[:24])
        self.assertEqual(task['task_type'],'solve')
        self.assertEqual(task['origin'],'shadow-ai')
        status=json.loads(await runner.run('project_task_status',json.dumps({'task_id':task['id']})))
        self.assertEqual(status['state'],'ready')
        self.assertIn("+print('new')",task['patch'])
        self.assertEqual((source/'app.py').read_text(),"print('old')\n")
        other=AgentRunner(self.assistant,Path(self.tmp.name)/'other',22,{},dev)
        self.assertIn('error',json.loads(await other.run('project_task_status',json.dumps({'task_id':task['id']}))))
        await dev.close()

    async def test_generated_files_require_owner_auth_and_are_not_public_paths(self):
        from shadow import app as module
        self.compat.chat.completions.create.side_effect=[completion(calls=[tool_call('a','create_file',{'name':'private.txt','content':'account-private'})]),completion('Tayyor')]
        job=self.worker.start(11,self.assistant,'Create a private file','agent')
        await self.worker.jobs[job['id']]['task']
        thread=self.worker.store.get(11,job['thread_id']);file=thread['messages'][-1]['attachments'][0]
        path=f"/dashboard/api/workspace/threads/{thread['id']}/files/{file['id']}"
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=module.app),base_url='http://test') as client:
            self.assertEqual((await client.get(path)).status_code,401)
            with patch.object(module,'_dashboard_allowed',return_value=True),patch.object(module,'workspace_ai',self.worker):
                with patch.object(module,'agent',SimpleNamespace(account_id=11)):
                    response=await client.get(path+'?account=11')
                    self.assertEqual(response.status_code,200)
                    self.assertEqual(response.text,'account-private')
                    self.assertEqual(response.headers['cache-control'],'no-store')
                    self.assertIn('private.txt',response.headers['content-disposition'])
                with patch.object(module,'agent',SimpleNamespace(account_id=22)):
                    self.assertEqual((await client.get(path+'?account=22')).status_code,404)
                    self.assertEqual((await client.get(path+'?account=11')).status_code,409)

    async def test_cancel_stops_this_commands_background_project_worker(self):
        dev=DevelopmentStudio(root=Path(self.tmp.name),state_dir=Path(self.tmp.name)/'dev',generate=AsyncMock())
        worker=Workspace(self.worker.store,development=dev)
        async def wait_dev(job,*args):
            await asyncio.Event().wait()
        async def start_then_wait(*args,runner,**kwargs):
            await runner.run('start_project_task','{"objective":"Fix this Shadow task","kind":"build"}')
            await asyncio.Event().wait()
        with patch.object(dev,'_run',side_effect=wait_dev),patch('shadow.workspace.workspace_reply',side_effect=start_then_wait):
            job=worker.start(11,self.assistant,'Fix this project','agent')
            await asyncio.sleep(.02)
            result=await worker.cancel(11,job['id'])
        self.assertEqual(result['state'],'cancelled')
        self.assertEqual(next(iter(dev.jobs.values()))['state'],'cancelled')
        await dev.close()
        await worker.close()

    async def test_model_loop_is_bounded_even_when_it_ignores_final_tool_choice(self):
        self.compat.chat.completions.create.return_value=completion(calls=[tool_call('a','calculate',{'expression':'2+2'})])
        job=self.worker.start(11,self.assistant,'Calculate then finish','agent')
        await self.worker.jobs[job['id']]['task']
        self.assertEqual(self.compat.chat.completions.create.await_count,6)
        self.assertEqual(self.worker.jobs[job['id']]['state'],'error')
        receipt=self.worker.store.get(11,job['thread_id'])['messages'][-1]
        self.assertEqual(len(receipt['events']),1)

    async def test_unsupported_tool_calling_does_not_silently_claim_execution(self):
        self.compat.chat.completions.create.side_effect=BadRequest()
        job=self.worker.start(11,self.assistant,'Create report file','agent')
        await self.worker.jobs[job['id']]['task']
        result=self.worker.public(self.worker.jobs[job['id']])
        self.assertEqual(result['state'],'error')
        self.assertIn('agent tools',result['error'])
        self.assertEqual(len(self.worker.store.get(11,job['thread_id'])['messages']),1)
