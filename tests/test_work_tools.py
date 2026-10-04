import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from shadow.office_files import OfficeFileError, inspect_office
from shadow.work_tools import WorkToolError, calculate, create_excel, create_word


class WorkToolsTests(unittest.TestCase):
    def test_arithmetic_and_rejected_code(self):
        self.assertEqual(calculate("(125 + 75) * 0.15"), "30.0")
        self.assertEqual(calculate("sqrt(144)"), "12.0")
        self.assertAlmostEqual(float(calculate("sin(pi/2)")), 1)
        for expression in ["__import__('os').system('id')", "(1).__class__", "2**10000", "[1, 2]"]:
            with self.subTest(expression=expression):
                with self.assertRaises((WorkToolError, ValueError)):
                    calculate(expression)

    def test_excel_round_trip_and_formulas(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            path = create_excel({"sheets": [{"name": "Budget", "rows": [
                ["Item", "Amount"], ["Salary", 1200], ["Total", "=SUM($B$2:$B$2)"],
            ]}]}, directory, 1)
            book = load_workbook(path)
            self.assertEqual(book.active["B3"].value, "=SUM($B$2:$B$2)")
            self.assertEqual(book.active["B2"].value, 1200)
            self.assertEqual(book.active.freeze_panes, "A2")
            book.close()
            preview = inspect_office(path)
            self.assertIn("Salary", preview)
            self.assertIn("=SUM($B$2:$B$2)", preview)
            for formula in ['=WEBSERVICE("https://example.com")', "='[secret.xlsx]Sheet1'!A1"]:
                with self.assertRaises(WorkToolError):
                    create_excel({"sheets": [{"name": "Test", "rows": [[formula]]}]}, directory, 2)

    def test_word_round_trip(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = create_word({
                "title": "Hisobot", "sections": [{
                    "heading": "Natijalar", "paragraphs": ["Jami 30 ta."],
                    "table": [["Nomi", "Soni"], ["A", 30]],
                }],
            }, Path(temporary), 1)
            doc = Document(path)
            self.assertEqual(doc.paragraphs[0].text, "Hisobot")
            self.assertEqual(doc.tables[0].cell(1, 1).text, "30")
            self.assertIn("Natijalar", inspect_office(path))

    def test_bad_zip_and_truncation(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            invalid = directory / "bad.xlsx"
            invalid.write_text("not a ZIP")
            with self.assertRaises(OfficeFileError):
                inspect_office(invalid)
            path = create_excel({"sheets": [{"name": "Large", "rows": [
                ["Row"], *[[index] for index in range(220)],
            ]}]}, directory, 1)
            self.assertIn("TRUNCATED", inspect_office(path))


if __name__ == "__main__":
    unittest.main()
